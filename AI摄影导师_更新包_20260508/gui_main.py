"""
AI摄影导师 - PySide6 GUI主界面
清新自然风格，实现图片上传、参数控制、分析执行、结果展示
"""
import sys, os, json, cv2
import numpy as np
import requests
from datetime import datetime
from PySide6.QtWidgets import *
from PySide6.QtCore import *
from PySide6.QtGui import *

from composition_analyzer import CompositionAnalyzer
from llm_evaluator import LLMEvaluator
from history_manager import HistoryManager
from gui_widgets import (Colors, GLOBAL_STYLE, ScoreRing, CardWidget,
                         StyledButton, SelectableImageLabel, SAMPLE_IMAGES)
from case_library import EXCELLENT_CASES, NEEDS_IMPROVEMENT_CASES, SPECIAL_CASES


class ImageDownloader(QThread):
    finished = Signal(str, str)
    error = Signal(str, str)
    def __init__(self, url, save_dir, parent=None):
        super().__init__(parent)
        self.url = url
        self.save_dir = save_dir
    def run(self):
        try:
            os.makedirs(self.save_dir, exist_ok=True)
            resp = requests.get(self.url, timeout=15, stream=True)
            if resp.status_code == 200:
                ext = self.url.split(".")[-1].split("?")[0][:4]
                if ext not in ("jpg","jpeg","png","webp","bmp"): ext = "jpg"
                fn = f"sample_{hash(self.url)%100000}.{ext}"
                path = os.path.join(self.save_dir, fn)
                with open(path,"wb") as f:
                    for chunk in resp.iter_content(8192): f.write(chunk)
                self.finished.emit(self.url, path)
            else: self.error.emit(self.url, f"HTTP {resp.status_code}")
        except Exception as e: self.error.emit(self.url, str(e))


class AnalysisWorker(QThread):
    progress = Signal(str)
    analysis_done = Signal(dict)
    comment_chunk = Signal(str)
    finished_all = Signal()
    error = Signal(str)
    def __init__(self, image_path, rules, style, roi=None, parent=None):
        super().__init__(parent)
        self.image_path = image_path
        self.rules = rules
        self.style = style
        self.roi = roi
    def run(self):
        try:
            self.progress.emit("正在进行构图分析...")
            analyzer = CompositionAnalyzer()
            result = analyzer.analyze(self.image_path, self.rules, self.roi)
            if "error" in result:
                self.error.emit(result["error"]); return
            self.analysis_done.emit(result)
            self.progress.emit("构图分析完成，正在生成AI评语...")
            evaluator = LLMEvaluator()
            for chunk in evaluator.generate_comment_stream(result, self.style):
                self.comment_chunk.emit(chunk)
            self.progress.emit("分析完成！")
            self.finished_all.emit()
        except Exception as e:
            self.error.emit(f"分析过程出错: {str(e)}")


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.image_path = None
        self.analysis_result = None
        self.overlay_image = None
        self.worker = None
        self._showing_overlay = False
        self._downloaders = []
        self._roi = None  # 交互式主体框选区域 (x, y, w, h)
        app_dir = os.path.dirname(os.path.abspath(__file__))
        self.sample_dir = os.path.join(app_dir, "data", "samples")
        self.history_mgr = HistoryManager()
        self._init_ui()
        self._check_ollama()
        self._load_history()

    def _init_ui(self):
        self.setWindowTitle("AI摄影导师 - 构图辅助分析系统")
        self.setMinimumSize(1280, 820)
        self.setStyleSheet(GLOBAL_STYLE)

        central = QWidget()
        self.setCentralWidget(central)
        root = QHBoxLayout(central)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        # ===== 左侧边栏 =====
        sidebar = QWidget()
        sidebar.setFixedWidth(380)
        sidebar.setStyleSheet(f"background-color: {Colors.SIDEBAR_BG};")
        sb_layout = QVBoxLayout(sidebar)
        sb_layout.setContentsMargins(18, 20, 18, 16)
        sb_layout.setSpacing(10)

        # Logo区
        logo_lbl = QLabel("🌿 AI摄影导师")
        logo_lbl.setStyleSheet(f"font-size: 20px; font-weight: 700; color: {Colors.PRIMARY}; padding: 4px 0 2px 0; background: transparent;")
        logo_lbl.setWordWrap(True)
        sb_layout.addWidget(logo_lbl)
        sub_lbl = QLabel("基于图像构图规则与LLM评语的构图辅助分析系统")
        sub_lbl.setWordWrap(True)
        sub_lbl.setStyleSheet(f"font-size: 11px; color: {Colors.TEXT_SEC}; background: transparent;")
        sb_layout.addWidget(sub_lbl)

        # 分隔线
        line = QFrame()
        line.setFrameShape(QFrame.HLine)
        line.setStyleSheet(f"background-color: {Colors.BORDER}; max-height: 1px;")
        sb_layout.addWidget(line)

        # 图片上传区
        self.image_label = QLabel("点击下方按钮\n选择一张图片")
        self.image_label.setAlignment(Qt.AlignCenter)
        self.image_label.setMinimumHeight(200)
        self.image_label.setStyleSheet(f"background-color: {Colors.CARD_BG}; border: 2px dashed {Colors.BORDER}; border-radius: 10px; color: {Colors.TEXT_LIGHT}; font-size: 14px;")
        sb_layout.addWidget(self.image_label)

        btn_row = QHBoxLayout()
        self.btn_upload = StyledButton("选择图片", "primary")
        self.btn_upload.clicked.connect(self._on_upload)
        btn_row.addWidget(self.btn_upload)
        self.btn_sample = StyledButton("示例图片", "outline")
        self.btn_sample.clicked.connect(self._show_sample_menu)
        btn_row.addWidget(self.btn_sample)
        sb_layout.addLayout(btn_row)

        # 参数区
        param_card = CardWidget("分析参数")
        param_card.setStyleSheet(f"CardWidget {{ background-color: {Colors.CARD_BG}; border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 12px; }}")
        rl = QLabel("构图规则")
        rl.setStyleSheet(f"font-size: 13px; color: {Colors.TEXT_SEC}; font-weight: 500; background: transparent;")
        param_card.add_widget(rl)
        self.chk_thirds = QCheckBox("三分法分析")
        self.chk_thirds.setChecked(True)
        self.chk_thirds.setStyleSheet(f"QCheckBox {{ font-size: 13px; background: transparent; }} QCheckBox::indicator:checked {{ background-color: {Colors.PRIMARY}; border: 1px solid {Colors.PRIMARY}; border-radius: 3px; }}")
        self.chk_symmetry = QCheckBox("对称性分析")
        self.chk_symmetry.setChecked(True)
        self.chk_symmetry.setStyleSheet(self.chk_thirds.styleSheet())
        param_card.add_widget(self.chk_thirds)
        param_card.add_widget(self.chk_symmetry)

        sl = QLabel("评语风格")
        sl.setStyleSheet(f"font-size: 13px; color: {Colors.TEXT_SEC}; font-weight: 500; margin-top: 4px; background: transparent;")
        param_card.add_widget(sl)
        self.combo_style = QComboBox()
        self.combo_style.addItems(["专业", "简洁", "鼓励"])
        self.combo_style.setMinimumWidth(120)
        self.combo_style.setStyleSheet(f"QComboBox {{ padding: 6px 10px; border: 1px solid {Colors.BORDER}; border-radius: 6px; background: white; font-size: 13px; }} QComboBox::drop-down {{ border: none; }} QComboBox QAbstractItemView {{ background: white; border: 1px solid {Colors.BORDER}; selection-background-color: {Colors.PRIMARY_LIGHT}; }}")
        param_card.add_widget(self.combo_style)
        sb_layout.addWidget(param_card)

        # 开始分析按钮
        self.btn_analyze = StyledButton("开始分析", "primary")
        self.btn_analyze.setMinimumHeight(46)
        self.btn_analyze.setEnabled(False)
        self.btn_analyze.clicked.connect(self._on_analyze)
        sb_layout.addWidget(self.btn_analyze)

        # 进度
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 0)
        self.progress_bar.setVisible(False)
        self.progress_bar.setMaximumHeight(4)
        self.progress_bar.setStyleSheet(f"QProgressBar {{ border: none; background: {Colors.BORDER_LIGHT}; border-radius: 2px; }} QProgressBar::chunk {{ background: {Colors.PRIMARY}; border-radius: 2px; }}")
        sb_layout.addWidget(self.progress_bar)

        self.lbl_status = QLabel("")
        self.lbl_status.setStyleSheet(f"color: {Colors.TEXT_SEC}; font-size: 12px; background: transparent;")
        sb_layout.addWidget(self.lbl_status)

        # Ollama状态
        self.lbl_ollama = QLabel("")
        self.lbl_ollama.setWordWrap(True)
        self.lbl_ollama.setStyleSheet(f"font-size: 11px; color: {Colors.TEXT_LIGHT}; background: transparent;")
        sb_layout.addWidget(self.lbl_ollama)

        sb_layout.addStretch()

        # 历史记录区
        hist_card = CardWidget("历史记录")
        hist_card.setStyleSheet(f"CardWidget {{ background-color: {Colors.CARD_BG}; border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 12px; }}")
        self.combo_history = QComboBox()
        self.combo_history.setPlaceholderText("暂无历史记录")
        self.combo_history.setStyleSheet(self.combo_style.styleSheet())
        self.combo_history.currentIndexChanged.connect(self._on_history_select)
        hist_card.add_widget(self.combo_history)
        hbtn = QHBoxLayout()
        self.btn_export = StyledButton("导出报告", "accent")
        self.btn_export.setEnabled(False)
        self.btn_export.clicked.connect(self._on_export)
        hbtn.addWidget(self.btn_export)
        self.btn_clear = StyledButton("清空", "default")
        self.btn_clear.clicked.connect(self._on_clear_history)
        hbtn.addWidget(self.btn_clear)
        hist_card.add_layout(hbtn)
        sb_layout.addWidget(hist_card)

        root.addWidget(sidebar)

        # ===== 右侧主内容区（标签页） =====
        right = QWidget()
        right.setStyleSheet(f"background-color: {Colors.BG};")
        right_layout = QVBoxLayout(right)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_layout.setSpacing(0)

        self.tab_widget = QTabWidget()
        self.tab_widget.setStyleSheet(f"""
            QTabWidget::pane {{ border: none; background: {Colors.BG}; }}
            QTabBar::tab {{ padding: 10px 28px; font-size: 14px; font-weight: 500; background: {Colors.SIDEBAR_BG}; border: none; border-bottom: 3px solid transparent; color: {Colors.TEXT_SEC}; }}
            QTabBar::tab:selected {{ background: {Colors.BG}; color: {Colors.PRIMARY}; border-bottom: 3px solid {Colors.PRIMARY}; }}
            QTabBar::tab:hover {{ color: {Colors.PRIMARY}; }}
        """)

        # ---- 标签页1：构图分析 ----
        analysis_page = QWidget()
        analysis_layout = QVBoxLayout(analysis_page)
        analysis_layout.setContentsMargins(20, 16, 20, 16)
        analysis_layout.setSpacing(16)

        # 顶部欢迎/标题
        self.header_label = QLabel("上传一张照片，开始AI构图分析")
        self.header_label.setStyleSheet(f"font-size: 20px; font-weight: 600; color: {Colors.TEXT}; background: transparent;")
        analysis_layout.addWidget(self.header_label)

        # 分数展示区
        score_row = QHBoxLayout()
        score_row.setSpacing(20)
        self.score_thirds = ScoreRing(0, "三分法", 100)
        self.score_symmetry = ScoreRing(0, "对称性", 100)
        score_row.addWidget(self.score_thirds)
        score_row.addWidget(self.score_symmetry)
        score_row.addStretch()

        self.btn_toggle = StyledButton("切换叠加层", "outline")
        self.btn_toggle.setEnabled(False)
        self.btn_toggle.clicked.connect(self._toggle_overlay)
        score_row.addWidget(self.btn_toggle, alignment=Qt.AlignBottom)
        analysis_layout.addLayout(score_row)

        # 可视化分析图（支持交互式主体框选）
        vis_card = CardWidget("可视化分析（可在图上拖拽框选主体区域）")
        self.overlay_label = SelectableImageLabel()
        self.overlay_label.setAlignment(Qt.AlignCenter)
        self.overlay_label.setMinimumHeight(400)
        self.overlay_label.setStyleSheet(f"background-color: {Colors.SIDEBAR_BG}; border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 8px; color: {Colors.TEXT_LIGHT}; font-size: 14px;")
        self.overlay_label.setText("分析后将在此显示构图叠加图\n可在图上拖拽框选主体区域进行聚焦分析")
        self.overlay_label.roi_selected.connect(self._on_roi_selected)
        vis_card.add_widget(self.overlay_label)

        roi_row = QHBoxLayout()
        self.lbl_roi = QLabel("未框选主体区域")
        self.lbl_roi.setStyleSheet(f"font-size: 12px; color: {Colors.TEXT_LIGHT}; background: transparent;")
        roi_row.addWidget(self.lbl_roi)
        roi_row.addStretch()
        self.btn_clear_roi = StyledButton("清除框选", "default")
        self.btn_clear_roi.setMaximumWidth(100)
        self.btn_clear_roi.setEnabled(False)
        self.btn_clear_roi.clicked.connect(self._on_clear_roi)
        roi_row.addWidget(self.btn_clear_roi)
        vis_card.add_layout(roi_row)

        analysis_layout.addWidget(vis_card)

        # 下方两栏：量化数据 + AI评语
        bottom_row = QHBoxLayout()
        bottom_row.setSpacing(16)

        data_card = CardWidget("量化特征数据")
        self.txt_data = QTextEdit()
        self.txt_data.setReadOnly(True)
        self.txt_data.setMaximumHeight(180)
        self.txt_data.setStyleSheet(f"font-family: 'Consolas','Courier New',monospace; font-size: 13px; background: {Colors.SIDEBAR_BG}; border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 6px; padding: 8px; color: {Colors.TEXT};")
        data_card.add_widget(self.txt_data)
        bottom_row.addWidget(data_card)

        comment_card = CardWidget("AI摄影导师评语")
        self.txt_comment = QTextEdit()
        self.txt_comment.setReadOnly(True)
        self.txt_comment.setMaximumHeight(180)
        self.txt_comment.setStyleSheet(f"font-size: 14px; line-height: 1.6; background: {Colors.SIDEBAR_BG}; border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 6px; padding: 8px; color: {Colors.TEXT};")
        comment_card.add_widget(self.txt_comment)
        bottom_row.addWidget(comment_card)

        analysis_layout.addLayout(bottom_row)

        self.tab_widget.addTab(analysis_page, "📷 构图分析")

        # ---- 标签页2：案例库 ----
        case_page = QWidget()
        case_layout = QVBoxLayout(case_page)
        case_layout.setContentsMargins(20, 16, 20, 16)
        case_layout.setSpacing(12)

        case_desc = QLabel("摄影构图案例库 — 包含构图优良、有待改进、特殊案例三类，每类20张参考图片")
        case_desc.setWordWrap(True)
        case_desc.setStyleSheet(f"font-size: 13px; color: {Colors.TEXT_SEC}; background: transparent; padding-bottom: 4px;")
        case_layout.addWidget(case_desc)

        self.case_tab = QTabWidget()
        self.case_tab.setStyleSheet(f"""
            QTabWidget::pane {{ border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 8px; background: {Colors.CARD_BG}; }}
            QTabBar::tab {{ padding: 8px 20px; font-size: 13px; background: {Colors.SIDEBAR_BG}; border: none; border-bottom: 2px solid transparent; color: {Colors.TEXT_SEC}; }}
            QTabBar::tab:selected {{ background: {Colors.CARD_BG}; color: {Colors.PRIMARY}; border-bottom: 2px solid {Colors.PRIMARY}; }}
        """)

        self.case_tab.addTab(self._build_case_grid(EXCELLENT_CASES, "excellent"), f"✅ 构图优良 ({len(EXCELLENT_CASES)})")
        self.case_tab.addTab(self._build_case_grid(NEEDS_IMPROVEMENT_CASES, "improve"), f"⚠️ 有待改进 ({len(NEEDS_IMPROVEMENT_CASES)})")
        self.case_tab.addTab(self._build_case_grid(SPECIAL_CASES, "special"), f"🎨 特殊案例 ({len(SPECIAL_CASES)})")
        case_layout.addWidget(self.case_tab)

        self.tab_widget.addTab(case_page, "📚 案例库")

        right_layout.addWidget(self.tab_widget)
        root.addWidget(right)

        self.statusBar().showMessage("就绪 - 请上传图片开始分析")
        self.statusBar().setStyleSheet(f"background: {Colors.SIDEBAR_BG}; color: {Colors.TEXT_SEC}; font-size: 12px; padding: 2px 8px;")

    # ========== 功能方法 ==========

    def _check_ollama(self):
        evaluator = LLMEvaluator()
        status = evaluator.check_connection()
        if status["connected"]:
            if status["target_model_available"]:
                self.lbl_ollama.setText("● Ollama已连接，Qwen2就绪")
                self.lbl_ollama.setStyleSheet(f"color: {Colors.SUCCESS}; font-size: 11px; background: transparent;")
            else:
                self.lbl_ollama.setText(f"● {status['message']}")
                self.lbl_ollama.setStyleSheet(f"color: {Colors.WARNING}; font-size: 11px; background: transparent;")
        else:
            self.lbl_ollama.setText(f"● {status['message']}")
            self.lbl_ollama.setStyleSheet(f"color: {Colors.ERROR}; font-size: 11px; background: transparent;")

    def _load_history(self):
        records = self.history_mgr.get_all_records()
        self.combo_history.blockSignals(True)
        self.combo_history.clear()
        for r in records:
            self.combo_history.addItem(f"[{r['timestamp']}] {r['filename']}", r['id'])
        self.combo_history.setCurrentIndex(-1)
        self.combo_history.blockSignals(False)

    def _show_sample_menu(self):
        menu = QMenu(self)
        menu.setStyleSheet(f"QMenu {{ background: white; border: 1px solid {Colors.BORDER}; border-radius: 6px; padding: 4px; }} QMenu::item {{ padding: 8px 16px; font-size: 13px; }} QMenu::item:selected {{ background: {Colors.PRIMARY_LIGHT}; color: {Colors.PRIMARY}; }}")
        for sample in SAMPLE_IMAGES:
            action = menu.addAction(f"{sample['name']} - {sample['desc']}")
            action.setData(sample)
        action = menu.exec(self.btn_sample.mapToGlobal(self.btn_sample.rect().bottomLeft()))
        if action:
            data = action.data()
            self._download_sample(data["url"], data["name"])

    def _download_sample(self, url, name):
        self.lbl_status.setText(f"正在下载示例图片: {name}...")
        self.btn_sample.setEnabled(False)
        dl = ImageDownloader(url, self.sample_dir)
        dl.finished.connect(self._on_sample_downloaded)
        dl.error.connect(self._on_sample_error)
        self._downloaders.append(dl)
        dl.start()

    def _on_sample_downloaded(self, url, path):
        self.btn_sample.setEnabled(True)
        self.lbl_status.setText("示例图片已加载")
        self.image_path = path
        self._display_image(path, self.image_label)
        self.btn_analyze.setEnabled(True)
        self.header_label.setText(f"已加载示例图片")
        self.statusBar().showMessage(f"已加载: {os.path.basename(path)}")

    def _on_sample_error(self, url, msg):
        self.btn_sample.setEnabled(True)
        self.lbl_status.setText(f"下载失败: {msg}")

    def _on_roi_selected(self, x, y, w, h):
        """用户在可视化区域框选了主体"""
        self._roi = (x, y, w, h)
        self.lbl_roi.setText(f"已框选主体区域: ({x}, {y}) {w}×{h}")
        self.lbl_roi.setStyleSheet(f"font-size: 12px; color: {Colors.PRIMARY}; font-weight: 500; background: transparent;")
        self.btn_clear_roi.setEnabled(True)

    def _on_clear_roi(self):
        """清除框选"""
        self._roi = None
        self.lbl_roi.setText("未框选主体区域")
        self.lbl_roi.setStyleSheet(f"font-size: 12px; color: {Colors.TEXT_LIGHT}; background: transparent;")
        self.btn_clear_roi.setEnabled(False)
        self.overlay_label.clear_roi()

    def _on_upload(self):
        path, _ = QFileDialog.getOpenFileName(self, "选择图片", "", "图片文件 (*.jpg *.jpeg *.png *.bmp *.tiff *.webp);;所有文件 (*)")
        if path:
            self.image_path = path
            self._display_image(path, self.image_label)
            self.btn_analyze.setEnabled(True)
            self.btn_toggle.setEnabled(False)
            self._showing_overlay = False
            self.overlay_label.setText("分析后将在此显示构图叠加图\n可在图上拖拽框选主体区域进行聚焦分析")
            self.overlay_label.setPixmap(QPixmap())
            self.overlay_label.clear_roi()
            self.txt_data.clear()
            self.txt_comment.clear()
            self.score_thirds.set_score(0)
            self.score_symmetry.set_score(0)
            self._roi = None
            self.lbl_roi.setText("未框选主体区域")
            self.lbl_roi.setStyleSheet(f"font-size: 12px; color: {Colors.TEXT_LIGHT}; background: transparent;")
            self.btn_clear_roi.setEnabled(False)
            self.header_label.setText(f"已加载: {os.path.basename(path)}")
            self.statusBar().showMessage(f"已加载图片: {os.path.basename(path)}")

    def _display_image(self, src, label):
        if isinstance(src, str):
            pixmap = QPixmap(src)
        elif isinstance(src, np.ndarray):
            img = src
            if len(img.shape) == 3:
                h, w, ch = img.shape
                rgb = cv2.cvtColor(img, cv2.COLOR_BGR2RGB)
                qimg = QImage(rgb.data, w, h, ch * w, QImage.Format_RGB888)
            else:
                h, w = img.shape
                qimg = QImage(img.data, w, h, w, QImage.Format_Grayscale8)
            pixmap = QPixmap.fromImage(qimg)
        else:
            return
        # 等比例缩放，确保图片完整显示
        label_w = max(label.width(), 400)
        label_h = max(label.height(), 300)
        scaled = pixmap.scaled(label_w, label_h, Qt.KeepAspectRatio, Qt.SmoothTransformation)
        label.setPixmap(scaled)
        label.setAlignment(Qt.AlignCenter)

    def _on_analyze(self):
        if not self.image_path: return
        rules = []
        if self.chk_thirds.isChecked(): rules.append("rule_of_thirds")
        if self.chk_symmetry.isChecked(): rules.append("symmetry")
        if not rules:
            QMessageBox.warning(self, "提示", "请至少选择一种构图规则"); return
        style = self.combo_style.currentText()
        self.btn_analyze.setEnabled(False)
        self.btn_upload.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.txt_data.clear()
        self.txt_comment.clear()
        self.header_label.setText("正在分析中...")
        self.worker = AnalysisWorker(self.image_path, rules, style, self._roi)
        self.worker.progress.connect(self._on_progress)
        self.worker.analysis_done.connect(self._on_analysis_done)
        self.worker.comment_chunk.connect(self._on_comment_chunk)
        self.worker.finished_all.connect(self._on_finished)
        self.worker.error.connect(self._on_error)
        self.worker.start()

    def _on_progress(self, msg):
        self.lbl_status.setText(msg)
        self.statusBar().showMessage(msg)

    def _on_analysis_done(self, result):
        self.analysis_result = result
        display = self._format_analysis_data(result)
        self.txt_data.setPlainText(display)
        analyses = result.get("analyses", {})
        if "rule_of_thirds" in analyses:
            self.score_thirds.set_score(analyses["rule_of_thirds"]["score"], "三分法")
        if "symmetry" in analyses:
            self.score_symmetry.set_score(analyses["symmetry"]["overall_score"], "对称性")
        analyzer = CompositionAnalyzer()
        overlay = analyzer.draw_analysis_overlay(self.image_path, result)
        if overlay is not None:
            self.overlay_image = overlay
            orig_img = cv2.imdecode(np.fromfile(self.image_path, dtype=np.uint8), cv2.IMREAD_COLOR)
            if orig_img is not None:
                h_o, w_o = orig_img.shape[:2]
                rgb_orig = cv2.cvtColor(orig_img, cv2.COLOR_BGR2RGB)
                qimg_orig = QImage(rgb_orig.data, w_o, h_o, 3 * w_o, QImage.Format_RGB888)
                self.overlay_label.set_image(QPixmap.fromImage(qimg_orig), w_o, h_o)
                h_ov, w_ov = overlay.shape[:2]
                rgb_ov = cv2.cvtColor(overlay, cv2.COLOR_BGR2RGB)
                qimg_ov = QImage(rgb_ov.data, w_ov, h_ov, 3 * w_ov, QImage.Format_RGB888)
                self.overlay_label.set_overlay(QPixmap.fromImage(qimg_ov))
                self.overlay_label.show_overlay(True)
                self._showing_overlay = True
            self.btn_toggle.setEnabled(True)

    def _on_comment_chunk(self, chunk):
        self.txt_comment.insertPlainText(chunk)
        cursor = self.txt_comment.textCursor()
        cursor.movePosition(cursor.MoveOperation.End)
        self.txt_comment.setTextCursor(cursor)

    def _on_finished(self):
        self.btn_analyze.setEnabled(True)
        self.btn_upload.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.btn_export.setEnabled(True)
        self.lbl_status.setText("分析完成")
        self.header_label.setText("分析完成")
        rules = []
        if self.chk_thirds.isChecked(): rules.append("rule_of_thirds")
        if self.chk_symmetry.isChecked(): rules.append("symmetry")
        rid = self.history_mgr.save_record(
            os.path.basename(self.image_path), self.image_path,
            self.analysis_result, self.txt_comment.toPlainText(),
            rules, self.combo_style.currentText()
        )
        self._load_history()

    def _on_error(self, msg):
        self.btn_analyze.setEnabled(True)
        self.btn_upload.setEnabled(True)
        self.progress_bar.setVisible(False)
        self.lbl_status.setText(f"错误: {msg}")
        self.header_label.setText("分析出错")
        QMessageBox.critical(self, "分析错误", msg)

    def _toggle_overlay(self):
        if self.overlay_image is not None:
            if self._showing_overlay:
                self.overlay_label.show_overlay(False)
                self._showing_overlay = False
                self.btn_toggle.setText("显示叠加层")
            else:
                self.overlay_label.show_overlay(True)
                self._showing_overlay = True
                self.btn_toggle.setText("显示原图")

    def _format_analysis_data(self, result):
        lines = []
        analyses = result.get("analyses", {})
        size = result.get("image_size", {})
        lines.append(f"图像尺寸: {size.get('width')} x {size.get('height')}")
        roi = result.get("roi")
        if roi:
            lines.append(f"主体框选: ({roi['x']}, {roi['y']}) {roi['width']}×{roi['height']}")
        lines.append("")
        if "rule_of_thirds" in analyses:
            rot = analyses["rule_of_thirds"]
            lines.append("━━━ 三分法分析 ━━━")
            lines.append(f"  得分: {rot['score']}/100")
            lines.append(f"  兴趣点数量: {len(rot.get('interest_points', []))}")
            lines.append(f"  平均归一化距离: {rot.get('avg_normalized_distance', 'N/A')}")
            lines.append(f"  判断: {rot.get('description', '')}")
            if rot.get('subject_center'):
                sc = rot['subject_center']
                lines.append(f"  主体中心: ({sc['x']}, {sc['y']})")
                lines.append(f"  主体到三分点距离: {rot.get('subject_to_thirds_distance', 'N/A')}")
            lines.append("")
        if "symmetry" in analyses:
            sym = analyses["symmetry"]
            lines.append("━━━ 对称性分析 ━━━")
            lines.append(f"  水平对称得分: {sym['horizontal_symmetry']['score']}/100")
            lines.append(f"  垂直对称得分: {sym['vertical_symmetry']['score']}/100")
            lines.append(f"  综合得分: {sym['overall_score']}/100")
            baxis = '水平' if sym['best_axis'] == 'horizontal' else '垂直'
            lines.append(f"  最佳对称轴: {baxis}")
            lines.append(f"  判断: {sym.get('description', '')}")
        if "visual_features" in analyses:
            vf = analyses["visual_features"]
            lines.append("")
            lines.append("━━━ 视觉特征 ━━━")
            b = vf.get("brightness", {})
            if b:
                lines.append(f"  亮度: {b.get('overall_mean', '')}/255 ({b.get('exposure_judgment', '')})")
                lines.append(f"  对比度: {b.get('contrast_level', '')} (标准差{b.get('overall_std', '')})")
            ln = vf.get("lines", {})
            if ln:
                lines.append(f"  线条: {ln.get('total_detected', 0)}条 ({ln.get('dominant_direction', '')})")
            vc = vf.get("visual_center", {})
            if vc:
                lines.append(f"  视觉重心: {vc.get('description', '')}")
            cp = vf.get("complexity", {})
            if cp:
                lines.append(f"  复杂度: {cp.get('edge_density_percent', '')}% ({cp.get('level', '')})")
            co = vf.get("color", {})
            if co:
                lines.append(f"  色调: {co.get('dominant_tone', '')} ({co.get('saturation_level', '')})")
        return "\n".join(lines)

    def _on_history_select(self, index):
        if index < 0: return
        rid = self.combo_history.itemData(index)
        if rid is None: return
        record = self.history_mgr.get_record_by_id(rid)
        if not record: return
        self.image_path = record["image_path"]
        if os.path.exists(record["image_path"]):
            self._display_image(record["image_path"], self.image_label)
        self.analysis_result = record["analysis"]
        self.txt_data.setPlainText(self._format_analysis_data(record["analysis"]))
        self.txt_comment.setPlainText(record["comment"])
        analyses = record["analysis"].get("analyses", {})
        if "rule_of_thirds" in analyses:
            self.score_thirds.set_score(analyses["rule_of_thirds"]["score"], "三分法")
        if "symmetry" in analyses:
            self.score_symmetry.set_score(analyses["symmetry"]["overall_score"], "对称性")
        if os.path.exists(record["image_path"]):
            analyzer = CompositionAnalyzer()
            overlay = analyzer.draw_analysis_overlay(record["image_path"], record["analysis"])
            if overlay is not None:
                self.overlay_image = overlay
                orig_img = cv2.imdecode(np.fromfile(record["image_path"], dtype=np.uint8), cv2.IMREAD_COLOR)
                if orig_img is not None:
                    h_o, w_o = orig_img.shape[:2]
                    rgb_orig = cv2.cvtColor(orig_img, cv2.COLOR_BGR2RGB)
                    qimg_orig = QImage(rgb_orig.data, w_o, h_o, 3 * w_o, QImage.Format_RGB888)
                    self.overlay_label.set_image(QPixmap.fromImage(qimg_orig), w_o, h_o)
                    h_ov, w_ov = overlay.shape[:2]
                    rgb_ov = cv2.cvtColor(overlay, cv2.COLOR_BGR2RGB)
                    qimg_ov = QImage(rgb_ov.data, w_ov, h_ov, 3 * w_ov, QImage.Format_RGB888)
                    self.overlay_label.set_overlay(QPixmap.fromImage(qimg_ov))
                    self.overlay_label.show_overlay(True)
                    self._showing_overlay = True
                self.btn_toggle.setEnabled(True)
        self.btn_export.setEnabled(True)
        self.header_label.setText(f"历史记录: {record['filename']}")

    def _on_export(self):
        if not self.analysis_result: return
        path, _ = QFileDialog.getSaveFileName(self, "导出报告",
            f"构图分析报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
            "Word文档 (*.docx);;PDF文件 (*.pdf);;文本文件 (*.txt);;JSON文件 (*.json);;所有文件 (*)")
        if not path: return
        try:
            if path.endswith(".json"):
                report = {"analysis": self.analysis_result, "comment": self.txt_comment.toPlainText(), "timestamp": datetime.now().isoformat()}
                with open(path, "w", encoding="utf-8") as f:
                    json.dump(report, f, ensure_ascii=False, indent=2)
            elif path.endswith(".docx"):
                self._export_word(path)
            elif path.endswith(".pdf"):
                self._export_pdf(path)
            else:
                with open(path, "w", encoding="utf-8") as f:
                    f.write("=" * 60 + "\n")
                    f.write("        AI摄影导师 - 构图分析报告\n")
                    f.write("=" * 60 + "\n\n")
                    f.write(f"分析时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                    f.write(f"图片文件: {os.path.basename(self.image_path)}\n\n")
                    f.write("--- 量化分析数据 ---\n")
                    f.write(self.txt_data.toPlainText())
                    f.write("\n\n--- AI评语 ---\n")
                    f.write(self.txt_comment.toPlainText())
                    f.write("\n")
            self.statusBar().showMessage(f"报告已导出: {path}")
            QMessageBox.information(self, "导出成功", f"报告已保存到:\n{path}")
        except Exception as e:
            QMessageBox.critical(self, "导出失败", str(e))

    def _export_word(self, path):
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        # 标题
        title = doc.add_heading("AI摄影导师 - 构图分析报告", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"分析时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph(f"图片文件：{os.path.basename(self.image_path)}")
        # 插入图片
        if self.image_path and os.path.exists(self.image_path):
            doc.add_heading("原始图片", level=2)
            doc.add_picture(self.image_path, width=Inches(5))
        # 量化数据
        doc.add_heading("量化分析数据", level=2)
        for line in self.txt_data.toPlainText().split("\n"):
            if line.strip():
                doc.add_paragraph(line)
        # AI评语
        doc.add_heading("AI摄影导师评语", level=2)
        doc.add_paragraph(self.txt_comment.toPlainText())
        doc.save(path)

    def _export_pdf(self, path):
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib.colors import HexColor
        # 注册中文字体
        font_path = None
        for fp in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc", "C:/Windows/Fonts/simhei.ttf"]:
            if os.path.exists(fp):
                font_path = fp; break
        if font_path:
            try:
                pdfmetrics.registerFont(TTFont("ChineseFont", font_path))
                cn_font = "ChineseFont"
            except:
                cn_font = "Helvetica"
        else:
            cn_font = "Helvetica"
        c = canvas.Canvas(path, pagesize=A4)
        w, h = A4
        y = h - 40 * mm
        # 标题
        c.setFont(cn_font, 18)
        c.setFillColor(HexColor("#5B8C5A"))
        c.drawCentredString(w / 2, y, "AI摄影导师 - 构图分析报告")
        y -= 12 * mm
        c.setFont(cn_font, 10)
        c.setFillColor(HexColor("#333333"))
        c.drawString(25 * mm, y, f"分析时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        y -= 6 * mm
        c.drawString(25 * mm, y, f"图片文件：{os.path.basename(self.image_path)}")
        y -= 10 * mm
        # 量化数据
        c.setFont(cn_font, 13)
        c.setFillColor(HexColor("#5B8C5A"))
        c.drawString(25 * mm, y, "量化分析数据")
        y -= 7 * mm
        c.setFont(cn_font, 9)
        c.setFillColor(HexColor("#333333"))
        for line in self.txt_data.toPlainText().split("\n"):
            if line.strip():
                c.drawString(28 * mm, y, line[:80])
                y -= 5 * mm
                if y < 30 * mm:
                    c.showPage(); y = h - 25 * mm
                    c.setFont(cn_font, 9); c.setFillColor(HexColor("#333333"))
        y -= 5 * mm
        # AI评语
        c.setFont(cn_font, 13)
        c.setFillColor(HexColor("#5B8C5A"))
        c.drawString(25 * mm, y, "AI摄影导师评语")
        y -= 7 * mm
        c.setFont(cn_font, 9)
        c.setFillColor(HexColor("#333333"))
        comment = self.txt_comment.toPlainText()
        # 简单自动换行
        max_chars = 55
        for line in comment.split("\n"):
            while len(line) > max_chars:
                c.drawString(28 * mm, y, line[:max_chars])
                line = line[max_chars:]
                y -= 5 * mm
                if y < 30 * mm:
                    c.showPage(); y = h - 25 * mm
                    c.setFont(cn_font, 9); c.setFillColor(HexColor("#333333"))
            c.drawString(28 * mm, y, line)
            y -= 5 * mm
            if y < 30 * mm:
                c.showPage(); y = h - 25 * mm
                c.setFont(cn_font, 9); c.setFillColor(HexColor("#333333"))
        c.save()

    def _build_case_grid(self, cases, category):
        """构建案例库网格视图"""
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea { border: none; background: transparent; }")
        container = QWidget()
        grid = QGridLayout(container)
        grid.setSpacing(12)
        grid.setContentsMargins(12, 12, 12, 12)
        cols = 4
        for i, case in enumerate(cases):
            card = QFrame()
            card.setStyleSheet(f"QFrame {{ background: {Colors.CARD_BG}; border: 1px solid {Colors.BORDER_LIGHT}; border-radius: 8px; }} QFrame:hover {{ border-color: {Colors.PRIMARY}; }}")
            card.setCursor(Qt.PointingHandCursor)
            cl = QVBoxLayout(card)
            cl.setContentsMargins(8, 8, 8, 8)
            cl.setSpacing(4)
            # 图片占位
            img_lbl = QLabel(f"📷 {case['name']}")
            img_lbl.setAlignment(Qt.AlignCenter)
            img_lbl.setMinimumHeight(100)
            img_lbl.setStyleSheet(f"background: {Colors.SIDEBAR_BG}; border-radius: 6px; font-size: 12px; color: {Colors.TEXT_SEC};")
            cl.addWidget(img_lbl)
            # 名称
            name_lbl = QLabel(case["name"])
            name_lbl.setStyleSheet(f"font-size: 12px; font-weight: 600; color: {Colors.TEXT}; background: transparent;")
            name_lbl.setWordWrap(True)
            cl.addWidget(name_lbl)
            # 描述
            desc_lbl = QLabel(case["desc"])
            desc_lbl.setStyleSheet(f"font-size: 11px; color: {Colors.TEXT_SEC}; background: transparent;")
            desc_lbl.setWordWrap(True)
            cl.addWidget(desc_lbl)
            # 使用按钮
            btn = StyledButton("加载分析", "outline")
            btn.setMaximumHeight(30)
            btn.setProperty("case_url", case["url"])
            btn.setProperty("case_name", case["name"])
            btn.clicked.connect(lambda checked, u=case["url"], n=case["name"]: self._load_case_image(u, n))
            cl.addWidget(btn)
            grid.addWidget(card, i // cols, i % cols)
        scroll.setWidget(container)
        return scroll

    def _load_case_image(self, url, name):
        """从案例库加载图片进行分析"""
        self.tab_widget.setCurrentIndex(0)  # 切换到分析页
        self._download_sample(url, name)

    def _on_clear_history(self):
        records = self.history_mgr.get_all_records()
        if records:
            reply = QMessageBox.question(self, "确认", "确定要清空所有历史记录吗？", QMessageBox.Yes | QMessageBox.No)
            if reply == QMessageBox.Yes:
                self.history_mgr.clear_all()
                self._load_history()
                self.btn_export.setEnabled(False)
